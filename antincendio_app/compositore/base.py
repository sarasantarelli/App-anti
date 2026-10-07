"""Funzioni di compilazione condivise tra i quattro template."""
from __future__ import annotations
import copy, re
from datetime import date
import docx
from docx.table import Table
from .. import docxkit as K
from ..agents import innesco as INN


def oggi() -> str:
    return date.today().strftime("%d/%m/%Y")


def nz(v, default=None):
    return default if v in (None, "", []) else v


def fmt_num(v, suffix=""):
    if v in (None, ""):
        return None
    f = float(v)
    s = f"{int(f)}" if f == int(f) else f"{f:g}".replace(".", ",")
    return s + suffix


class Base:
    def __init__(self, path, caso, voci=None, risposte=None, tieni_guide=False):
        self.doc = docx.Document(str(path))
        self.tabs = list(self.doc.tables)   # snapshot: gli indici delle voci si riferiscono al template originale
        self.c = caso
        self.voci = voci or []
        self.risp = risposte or {}
        self.tieni_guide = tieni_guide
        self.mancanti: list[str] = []

    # ---- dati comuni -------------------------------------------------------------
    def d(self, k, default=None):
        return self.c.get(k, default)

    def occ(self):
        return self.c.esito["normativo"]["occupanti"]

    def revisione_str(self):
        return f"Rev. {nz(self.d('revisione'), '00')} del {nz(self.d('data_emissione'), oggi())}"

    # ---- cover ------------------------------------------------------------------------
    def cover(self):
        for p in self.doc.paragraphs[:8]:
            t = p.text.strip()
            if t.startswith("[RAGIONE SOCIALE") or t in ("NOME AZIENDA",):
                if self.d("ragione_sociale"):
                    K.set_par(p, str(self.d("ragione_sociale")).upper())
            elif t.startswith("[Indirizzo") or t.startswith("Idirizzo"):
                if self.d("indirizzo"):
                    K.set_par(p, str(self.d("indirizzo")))

    # ---- tabelle etichetta|valore -------------------------------------------------------
    def label_fill(self, header_startswith: str, mapping: dict[str, str | None], occurrence=0):
        tbls = K.find_tables(self.doc, header_startswith)
        if len(tbls) <= occurrence:
            return
        _, t = tbls[occurrence]
        for row in t.rows:
            cs = K.uniq_cells(row)
            if len(cs) < 2:
                continue
            lab = K.ctext(cs[0])
            for key, val in mapping.items():
                if lab.lower().startswith(key.lower()) and val not in (None, ""):
                    cur = K.ctext(cs[1])
                    # conserva suffissi come '— da rilievo/planimetria'
                    m = re.match(r"^(?:\[[^\]]*\]|_+[^—]*)(\s*—.*)$", cur)
                    K.set_cell(cs[1], str(val) + (m.group(1) if m else ""))
                    break

    def anagrafica(self):
        sup = fmt_num(self.d("superficie_mq"), " m²")
        occ = self.occ()
        qmin, qmax = self.d("quota_min"), self.d("quota_max")
        quote = f"da {fmt_num(qmin)} m a {fmt_num(qmax)} m" if qmin is not None and qmax is not None else None
        self.label_fill("ANAGRAFICA", {
            "Ragione sociale": self.d("ragione_sociale"), "Sede": self.d("indirizzo"), "Codice ATECO": self.d("ateco"),
            "Superficie lorda": sup, "Affollamento": f"{int(occ)} occupanti" if occ else None, "Quote dei piani": quote})
        self.label_fill("FIGURE DI RIFERIMENTO", {"Datore di lavoro": self.d("datore_lavoro"), "RSPP": self.d("rspp"), "RLS": self.d("rls")})
        self.label_fill("DOCUMENTO", {"Data di emissione": nz(self.d("data_emissione"), oggi()), "Revisione": self.revisione_str()})
        self.label_fill("DVR GENERALE", {"Data ultima revisione DVR": self.d("dvr_data"), "Sezioni del DVR": self.d("dvr_sezioni"),
                                         "Piano di Emergenza collegato": self.d("pe_rif")})
        # registro revisioni
        for _, t in K.find_tables(self.doc, "Rev."):
            hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
            if len(hdr) >= 4 and hdr[1] == "Data":
                cs = K.uniq_cells(t.rows[1])
                K.set_cell(cs[0], str(nz(self.d("revisione"), "00")).zfill(2))
                K.set_cell(cs[1], nz(self.d("data_emissione"), oggi()))
                if self.d("redatto_da") or self.d("rspp"):
                    K.set_cell(cs[2], str(nz(self.d("redatto_da"), self.d("rspp"))))
        # sottoscrizioni di copertina
        names = {"DATORE": self.d("datore_lavoro"), "RSPP": self.d("rspp"), "RLS": self.d("rls")}
        for _, t in K.find_tables(self.doc, "SOTTOSCRIZIONI"):
            for row in t.rows:
                for cell in K.uniq_cells(row):
                    tx = K.ctext(cell)
                    for key, nm in names.items():
                        if nm and tx.upper().startswith(key) and "[Nome" in tx:
                            K.fill_tokens(cell, [str(nm)])

    def sottoscrizioni(self):
        names = (self.d("datore_lavoro"), self.d("rspp"), self.d("rls"))
        for t in self.doc.tables:
            if len(t.rows) == 2 and len(K.uniq_cells(t.rows[0])) in (2, 3):
                h = [K.ctext(x).lower() for x in K.uniq_cells(t.rows[0])]
                if h and h[0].startswith("datore di lavoro"):
                    for cell, nm in zip(K.uniq_cells(t.rows[1]), names):
                        if nm and "[Nome" in K.ctext(cell):
                            K.fill_tokens(cell, [str(nm)])

    # ---- paragrafi con segnaposto ---------------------------------------------------------
    def par_fill(self, startswith: str, values: list[str | None]):
        for p in self.doc.paragraphs:
            if p.text.strip().startswith(startswith):
                K.fill_tokens(p, values)
                return True
        return False

    def par_replace(self, startswith: str, new_text: str):
        for p in self.doc.paragraphs:
            if p.text.strip().startswith(startswith):
                K.set_par(p, new_text)
                return True
        return False

    # ---- attività ---------------------------------------------------------------------------
    def attivita(self):
        for _, t in K.find_tables(self.doc, "Aspetto"):
            m = {"Attività svolta": self.d("attivita_descrizione"), "Lavorazioni con rischio": self.d("lavorazioni_rischio"),
                 "Orari": self.d("orari"), "Presenza fuori orario": self.d("fuori_orario"), "Variabilità": self.d("variabilita")}
            for row in t.rows[1:]:
                cs = K.uniq_cells(row)
                for k, v in m.items():
                    if K.ctext(cs[0]).startswith(k) and v:
                        K.set_cell(cs[1], str(v))
        if self.d("confini"):
            self.par_fill("[Compartimentazioni", [str(self.d("confini"))])

    def locali(self):
        locali = self.d("locali") or []
        if not locali:
            locali = [{"nome": "Intera attività", "mq": self.d("superficie_mq"), "dest": ", ".join(self.d("tipologia_attivita", []) or []) or None}]
        for _, t in K.find_tables(self.doc, "Locale/Area"):
            hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
            if len(hdr) == 3 and hdr[1].startswith("Superficie"):
                K.ensure_rows(t, 1, len(locali))
                for i, l in enumerate(locali, start=1):
                    cs = K.uniq_cells(t.rows[i])
                    K.set_cell(cs[0], str(l.get("nome") or "—"))
                    K.set_cell(cs[1], fmt_num(l.get("mq"), " m²") or "[__ m²]")
                    K.set_cell(cs[2], str(l.get("dest") or "[destinazione]"))
                for j in range(len(t.rows) - 1, len(locali), -1):
                    if K.has_pending(t.rows[j].cells[0]):
                        K.delete_row(t, j)
                break

    def affollamento(self):
        vals = [self.d("occupanti_lavoratori"), self.d("occupanti_appaltatori"), self.d("occupanti_terzi"), self.d("occupanti_picco")]
        for _, t in K.find_tables(self.doc, "Categoria occupanti"):
            for i, v in enumerate(vals, start=1):
                if i < len(t.rows):
                    cs = K.uniq_cells(t.rows[i])
                    K.set_cell(cs[1], fmt_num(v) if v not in (None, "") else "0" if i < 4 else "—")
                    # familiarità (Codice): lavoratori familiari, pubblico non familiare
                    if len(cs) > 2 and "familiari" in K.ctext(cs[2]):
                        K.choose(cs[2], "familiari" if i in (1, 2) else "non familiari")
                    if len(cs) > 2 and K.ctext(cs[2]) in ("[]", ""):
                        K.set_cell(cs[2], "")
        tot = self.occ()
        self.par_fill("Totale occupanti", [f"{int(tot)}"])

    # ---- sostanze per area --------------------------------------------------------------------
    def sostanze(self):
        det = list(self.d("sostanze_dettaglio") or [])
        if not det:
            for s in self.d("sostanze", []) or []:
                det.append({"area": "Intera attività", "nome": s, "quantita": None, "unita": "", "stoccaggio": None, "modalita": None, "rilevanza": None})
        aree: dict[str, list] = {}
        for s in det:
            aree.setdefault(s.get("area") or "Intera attività", []).append(s)
        if not aree:
            aree = {"Intera attività": []}
        # coppie (titolo AREA n, tabella)
        bl = [x for x in K.blocks(self.doc) if isinstance(x, Table) or x.text.strip()]
        pairs = []
        for i, b in enumerate(bl[:-1]):
            if isinstance(b, Table) and len(b.rows) == 1 and K.ctext(b.rows[0].cells[0]).startswith("AREA ") and isinstance(bl[i + 1], Table):
                pairs.append((b, bl[i + 1]))
        if not pairs:
            return
        names = list(aree)
        # duplica blocchi se servono più aree
        while len(pairs) < len(names):
            tt, dt = pairs[-1]
            n_tt, n_dt = copy.deepcopy(tt._tbl), copy.deepcopy(dt._tbl)
            dt._tbl.addnext(n_tt); n_tt.addnext(n_dt)
            pairs.append((Table(n_tt, tt._parent), Table(n_dt, dt._parent)))
        for idx, (tt, dt) in enumerate(pairs):
            if idx >= len(names):
                K.remove_block(tt); K.remove_block(dt); continue
            nm = names[idx]
            K.set_cell(tt.rows[0].cells[0], f"AREA {idx + 1}: {nm} (da VRI-4)")
            rows = aree[nm] or []
            K.ensure_rows(dt, 1, max(1, len(rows)))
            ncol = len(K.uniq_cells(dt.rows[0]))
            for i, s in enumerate(rows, start=1):
                cs = K.uniq_cells(dt.rows[i])
                q = None
                if s.get("quantita") not in (None, ""):
                    q = f"{s['quantita']} {s.get('unita') or ''}".strip()
                for cell, val in zip(cs[:4], (s.get("nome"), q, s.get("stoccaggio"), s.get("modalita"))):
                    if val:
                        K.set_cell(cell, str(val))
                if s.get("rilevanza"):
                    K.choose(cs[4], str(s["rilevanza"]).capitalize())
            for j in range(len(dt.rows) - 1, max(len(rows), 1), -1):
                if K.has_pending(dt.rows[j].cells[0]):
                    K.delete_row(dt, j)

    # ---- qf per area (Minicodice VRI-5.1) -----------------------------------------------------
    def qf_area(self):
        r = self.c.esito["rischio"]
        qf = r.get("qf")
        for _, t in K.find_tables(self.doc, "Locale/Area"):
            hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
            if len(hdr) == 3 and hdr[1].startswith("Metodo"):
                cs = K.uniq_cells(t.rows[1])
                K.set_cell(cs[0], "Intera attività (area di riferimento)")
                K.set_cell(cs[1], r.get("qf_fonte", ""))
                K.set_cell(cs[2], f"{qf:,.0f} MJ/m²".replace(",", ".") if qf is not None else "[valore]")
        for _, t in K.find_tables(self.doc, "VERIFICA SOGLIA MINICODICE"):
            c0 = t.rows[0].cells[0]
            if qf is not None:
                K.fill_tokens(c0, [f"{qf:,.0f} MJ/m²".replace(",", ".")])
                K.choose(c0, "≤ 900 MJ/m²" if qf <= 900 else "> 900 MJ/m²")

    # ---- sorgenti di innesco ------------------------------------------------------------------
    def innesco(self):
        ev = INN.valuta(self.c)
        for _, t in K.find_tables(self.doc, "Categoria (UNI EN 1127-1)"):
            hdr = [K.ctext(x).lower() for x in K.uniq_cells(t.rows[0])]
            col = {"alta": 1, "media": 2, "bassa": 3}
            if "rara" in hdr:   # Minicodice: Freq./cont., Rara, Molto rara
                col = {"alta": 1, "media": 2, "bassa": 3}
            for i, e in enumerate(ev, start=1):
                if i >= len(t.rows):
                    break
                cs = K.uniq_cells(t.rows[i])
                if e["livello"]:
                    K.set_cell(cs[col[e["livello"]]], "☒")
                    K.set_cell(cs[-1], f"Riscontro da documentazione: {e['nota']}. Mitigazione e conferma in sopralluogo.")
                else:
                    K.set_cell(cs[-1], "N.A. — non rilevata nella documentazione esaminata (da confermare in sopralluogo).")

    # ---- check-list -------------------------------------------------------------------------------
    def checklist(self):
        for v in self.voci:
            r = self.risp.get(v.id)
            t = self.tabs[v.tab] if v.tab < len(self.tabs) else None
            if t is None or v.riga >= len(t.rows):
                continue
            cs = K.uniq_cells(t.rows[v.riga])
            if not r:
                continue
            es = r.get("esito")
            nota = r.get("nota") or ""
            if es and not nota:
                nota = {"C": "Conforme: verificato in sopralluogo", "NC": "Non conforme: riscontrato in sopralluogo",
                        "NA": "Non applicabile alla presente attività"}[es]
            if v.fmt == "CNC":
                if es:
                    K.choose(cs[3], {"C": "C", "NC": "NC", "NA": "N.A."}[es])
                    K.shade(cs[3], {"C": K.OK_FILL, "NC": K.NC_FILL}.get(es, "F0F0F0"))
                if nota:
                    K.set_cell(cs[2], nota)
                if es == "NC":
                    K.set_cell(cs[4], (r.get("azione") or "Ripristinare la conformità") + " — vedi VRI-11")
                elif es in ("C", "NA"):
                    K.set_cell(cs[4], "—")
            elif v.fmt == "RACC":
                if es:
                    K.choose(cs[1], {"C": "Sì", "NC": "No", "NA": "N.A."}[es])
                    K.choose(cs[3], {"C": "C", "NC": "NC", "NA": "N.A."}[es])
                    K.shade(cs[3], {"C": K.OK_FILL, "NC": K.NC_FILL}.get(es, "F0F0F0"))
                if nota:
                    K.set_cell(cs[2], nota)
            else:   # SNA
                if es:
                    for col, key in ((1, "C"), (2, "NC"), (3, "NA")):
                        K.set_cell(cs[col], "☒" if key == es else "☐")
                    if es == "NC":
                        K.shade(cs[2], K.NC_FILL)
                    elif es == "C":
                        K.shade(cs[1], K.OK_FILL)
                if nota:
                    K.set_cell(cs[4], nota)

    def esito_misura(self, misura: str):
        es = [self.risp.get(v.id, {}).get("esito") for v in self.voci if v.misura == misura]
        if not es:
            return None
        if "NC" in es:
            return "NC"
        if all(e in ("C", "NA") for e in es):
            return "C"
        return None

    # ---- azioni correttive ---------------------------------------------------------------------------
    def azioni(self):
        az = self.c.esito.get("azioni", [])
        for _, t in K.find_tables(self.doc, "Prior."):
            if not az and self.c.esito.get("completezza", {}).get("percento", 0) != 100:
                return   # sopralluogo incompleto: il piano azioni resta da compilare, non si dichiara «nessuna NC»
            if not az:
                for j in range(len(t.rows) - 1, 1, -1):
                    K.delete_row(t, j)
                cs = K.uniq_cells(t.rows[1])
                K.set_cell(cs[0], "—"); K.set_cell(cs[1], "Nessuna non conformità rilevata alla data di redazione; verifiche da ripetere alle scadenze previste.")
                for x in cs[2:]:
                    K.set_cell(x, "—")
                return
            K.ensure_rows(t, 1, len(az))
            for i, a in enumerate(az, start=1):
                cs = K.uniq_cells(t.rows[i])
                K.set_cell(cs[0], str(a["priorita"]))
                K.set_cell(cs[1], a["azione"])
                K.set_cell(cs[2], a.get("misura") or "—")
                K.set_cell(cs[3], a.get("responsabile") or "Datore di lavoro")
                K.set_cell(cs[4], a.get("scadenza") or {1: "entro 30 gg", 2: "entro 90 gg", 3: "entro 12 mesi", 4: "continuo"}[a["priorita"]])
                K.choose(cs[5], a.get("stato") or "Aperta")
                K.shade(cs[0], K.NC_FILL if a["priorita"] == 1 else K.PENDING_FILL)
            for j in range(len(t.rows) - 1, len(az), -1):
                K.delete_row(t, j)

    def periodicita(self):
        p = self.d("periodicita_revisione")
        self.par_fill("Periodicità di revisione ordinaria", [f"{int(float(p))} anni" if p else "3 anni"])

    def header_footer(self):
        nome = self.d("ragione_sociale")
        if not nome:
            return
        for sec in self.doc.sections:
            for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
                try:
                    paras = list(part.paragraphs) + [p for t in part.tables for r in t.rows for c in r.cells for p in c.paragraphs]
                except Exception:
                    continue
                for p in paras:
                    for r in p.runs:
                        if "[Ragione Sociale S.r.l.]" in r.text or "[RAGIONE SOCIALE S.r.l.]" in r.text:
                            r.text = r.text.replace("[Ragione Sociale S.r.l.]", str(nome)).replace("[RAGIONE SOCIALE S.r.l.]", str(nome))
                        elif r.text.strip() == "AZIENDA":
                            r.text = str(nome)

    def finalize(self):
        self.header_footer()
        K.strip_guides(self.doc, keep=self.tieni_guide)
        n = K.mark_pending(self.doc)
        return n

    def salva(self, out):
        self.doc.save(str(out))
