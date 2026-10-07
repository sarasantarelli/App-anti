"""Compila il template VRI Codice integrale (RTO residuale, DM 3/8/2015 sez. G e S)."""
from __future__ import annotations
import re
from pathlib import Path
from docx.table import Table
from .. import docxkit as K
from ..criteri import parse_criteri_tabelle, valuta, Ctx
from ..agents import strategia as ST
from .base import Base, nz, fmt_num

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "VRI_Codice_Integrale_RTO_completa_2.docx"
IT = lambda x, n=2: f"{x:.{n}f}".replace(".", ",")


def _ctx(caso):
    return ST.contesto(caso)


def compila(caso, voci, risposte, out, tieni_guide=False):
    b = Base(TEMPLATE, caso, voci, risposte, tieni_guide)
    n, r, dim = caso.esito["normativo"], caso.esito["rischio"], caso.esito["dimensionamento"]
    stra = caso.esito["strategia"]["misure"]
    ctx = _ctx(caso)
    b.cover(); b.anagrafica()

    # ---- VRI-2.1 percorso ----------------------------------------------------------------
    req = {x["id"]: x for x in n["requisiti_allegato_I"]}
    scr = n["screening"]
    for _, t in K.find_tables(b.doc, "Verifica"):
        if len(t.rows) == 11 and K.ctext(t.rows[1].cells[0]).startswith("A."):
            rows = {K.ctext(K.uniq_cells(rw)[0]).split(" ")[0]: K.uniq_cells(rw) for rw in t.rows[1:]}
            soggetta = n["soggetta_dpr151"]
            cs = rows["A."]
            K.choose(cs[2], "Sì" if soggetta else "No")
            K.set_cell(cs[1], ("voce n. " + ", ".join(str(x["voce"]) for x in scr if x["esito"] == "SOPRA")) if soggetta and scr else
                       "nessuna voce pertinente sopra soglia — " + n["fonte_assoggettamento"])
            K.choose(rows["A.1"][2], "N.A." if not soggetta else ("Sì" if b.d("rtv_applicabile") else "No"))
            K.choose(rows["B."][2], "Sì" if b.d("rt_settore") else "No")
            if b.d("rt_settore"):
                K.set_cell(rows["B."][1], str(b.d("rt_settore")))
            for key in ("C.a", "C.b", "C.c", "C.d", "C.e", "C.f", "D."):
                q = req["D" if key == "D." else key]
                cs = rows[key]
                if q["ok"] is True:
                    K.choose(cs[2], "Soddisfatto")
                elif q["ok"] is False:
                    K.choose(cs[2], "NON soddisfatto"); K.shade(cs[2], K.NC_FILL)
                K.set_cell(cs[1], f"{q['soglia']} — rilevato: {q['valore']}")
    ko = [f"{q['id']} {q['testo']} ({q['valore']}, soglia {q['soglia']})" for q in n["requisiti_allegato_I"] if q["ok"] is not True and q["id"] not in ("RTV",)]
    b.par_fill("[Requisito/i non soddisfatto/i]", ["Requisiti non soddisfatti: " + "; ".join(ko) if ko else "nessuno"])
    b.par_replace("[Requisito/i non soddisfatto/i]", "Requisiti non soddisfatti: " + ("; ".join(ko) if ko else "nessuno (verificare dati mancanti)"))
    # ---- VRI-4/5 ---------------------------------------------------------------------------
    b.attivita(); b.locali(); b.affollamento(); b.sostanze()

    # ---- VRI-6 aree a rischio specifico ---------------------------------------------------------
    sost = set(b.d("sostanze", []) or [])
    crit = []
    if b.d("sostanze_significative") or sost & {"liquidi infiammabili", "gpl/gas", "esplosivi/ATEX"}:
        crit.append("a")
    if b.d("lavorazioni_pericolose") or "lavorazioni a caldo" in sost:
        crit.append("b")
    if "olio/friggitrici" in sost or b.d("cucina_cottura"):
        crit.append("f")
    if "batterie/litio" in sost:
        crit.append("g")
    if r["qf"] and r["qf"] > 1200:
        crit.append("d")
    if r["rambiente_significativo"]:
        crit.append("h")
    rs_presente = bool(crit)
    for _, t in K.find_tables(b.doc, "Locale/Area"):
        hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
        if len(hdr) == 3 and hdr[1].startswith("Criterio"):
            cs = K.uniq_cells(t.rows[1])
            K.set_cell(cs[0], "Intera attività" if not b.d("locali") else "Aree individuate in VRI-4.2")
            if rs_presente:
                K.set_cell(cs[1], ", ".join(crit) + " — da confermare per singola area in sopralluogo (V.1.1 c.2)")
                K.choose(cs[2], "Area a rischio specifico")
            else:
                K.set_cell(cs[1], "Nessuno dei criteri a–h soddisfatto sui dati disponibili")
                K.choose(cs[2], "Nessun rischio specifico")
    b.innesco()

    # ---- VRI-7bis scenari ----------------------------------------------------------------------------
    inn = caso.esito.get("innesco", [])
    rank = {"alta": 3, "media": 2, "bassa": 1, None: 0}
    top = max(inn, key=lambda e: rank[e["livello"]]) if inn else None
    mat = (b.d("sostanze_dettaglio") or [{}])[0].get("nome") or (", ".join(sorted(sost)) if sost else None)
    for _, t in K.find_tables(b.doc, "Parametro dello scenario"):
        lab = {K.ctext(K.uniq_cells(rw)[0]): K.uniq_cells(rw) for rw in t.rows[1:]}
        if "Area/locale di riferimento" in lab:
            vals = {
                "Area/locale di riferimento": "Intera attività (locale con maggiore combinazione qf × affollamento da confermare in sopralluogo)",
                "Materiale/combustibile principale": mat or None,
                "Sorgente di innesco credibile": f"{top['categoria']} — {top['nota']}" if top and top["livello"] else None,
                "Velocità di crescita": f"δα = {r['dalfa']} — {r['dalfa_motivo']}",
                "Conseguenza attesa sugli occupanti": f"Rvita = {r['rvita']} (δocc {r['docc']} × δα {r['dalfa']})",
                "Conseguenza attesa sui beni": f"Rbeni = {r['rbeni']}",
                "Conseguenza attesa sull": "Rambiente " + ("significativo" if r["rambiente_significativo"] else "non significativo"),
            }
            for k, v in vals.items():
                for lk, cs in lab.items():
                    if lk.startswith(k) and v:
                        K.set_cell(cs[1], v)
            for lk, cs in lab.items():
                if lk.startswith("Condizioni che favoriscono"):
                    K.set_cell(cs[1], "Da definire in sopralluogo: compartimentazione, aperture non protette, accumuli di materiale combustibile")
        elif "Area a rischio specifico (da VRI-6)" in lab:
            if not rs_presente:
                for lk, cs in lab.items():
                    if lk.startswith("CONCLUSIONE"):
                        continue
                    K.set_cell(cs[1], "Non applicabile: nessuna area a rischio specifico individuata in VRI-6")
            for rw in t.rows:
                cell = K.uniq_cells(rw)[0]
                if K.ctext(cell).startswith("CONCLUSIONE SCENARI"):
                    K.fill_tokens(cell, [f"scenario convenzionale per aree ordinarie con incendio del materiale prevalente, δα = {r['dalfa']}, Rvita {r['rvita']}"])

    # ---- VRI-8 profili ---------------------------------------------------------------------------
    b.par_fill("δocc adottato per ciascun locale", [f"{r['docc']} — {r['docc_motivo']}"])
    b.par_fill("δα adottato", [f"δα = {r['dalfa']} — {r['dalfa_motivo']}"])
    for _, t in K.find_tables(b.doc, "δocc"):
        if len(t.rows) == 9:
            for rw in t.rows[1:]:
                for cell in K.uniq_cells(rw)[2:]:
                    if K.ctext(cell).split(" ")[0] == r["rvita"]:
                        K.set_cell(cell, f"{r['rvita']} ◄ adottato", bold=True)
    for _, t in K.find_tables(b.doc, "Locale/Area"):
        hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
        if len(hdr) == 2 and hdr[1].startswith("Rvita adottato"):
            rl = r.get("rvita_locali") or []
            if rl:
                K.ensure_rows(t, 1, len(rl) + 1)
                for i, x in enumerate(rl, start=1):
                    cs = K.uniq_cells(t.rows[i])
                    K.set_cell(cs[0], x["nome"]); K.set_cell(cs[1], f"{x['rvita']} — {x['motivo']}")
                cs = K.uniq_cells(t.rows[len(rl) + 1])
                K.set_cell(cs[0], "ADOTTATO PER LA STRATEGIA (profilo più elevato)", bold=True)
                K.set_cell(cs[1], f"{r['rvita']} — {r['docc_motivo']}; δα {r['dalfa']} ({r['dalfa_motivo']})", bold=True)
                continue
            cs = K.uniq_cells(t.rows[1])
            K.set_cell(cs[0], "Intera attività")
            K.set_cell(cs[1], f"{r['rvita']} — δocc {r['docc']} ({r['docc_motivo']}); δα {r['dalfa']} ({r['dalfa_motivo']})")
    for p in b.doc.paragraphs:
        if p.text.startswith("Vincolato:"):
            K.tick(p, "Sì" if r["vincolata"] else "No", nth=0)
            K.tick(p, "Sì" if r["strategica"] else "No", nth=1)
            K.fill_tokens(p, [str(r["rbeni"])])
    for _, t in K.find_tables(b.doc, "Condizione (Cap. G.3.4)"):
        for i, (lab, v) in enumerate(r["rambiente_condizioni"], start=1):
            K.choose(K.uniq_cells(t.rows[i])[1], "Sì" if v else "No")
    for p in b.doc.paragraphs:
        if p.text.startswith("Rambiente adottato"):
            K.choose(p, "Significativo (almeno" if r["rambiente_significativo"] else "Non significativo (nessuna")
            K.fill_tokens(p, ["condizioni valutate in base ai dati dell'attività (da confermare)"])

    # ---- carico d'incendio qf,d (scheda S.2.9) ---------------------------------------------------------
    q = r.get("qfd")
    if q:
        for _, t in K.find_tables(b.doc, "Parametro"):
            if len(t.rows) == 2 and K.ctext(t.rows[0].cells[0]) == "Parametro" and K.ctext(t.rows[0].cells[1]) == "Fonte":
                K.fill_tokens(t.rows[1].cells[2], [IT(q["qf"], 0)])
        for _, t in K.find_tables(b.doc, "Compartimento di riferimento"):
            cs = K.uniq_cells(t.rows[1])
            K.set_cell(cs[0], "Intera attività (compartimento di riferimento)")
            K.fill_tokens(cs[1], [IT(q["sup_comp"], 0)])
        for _, t in K.find_tables(b.doc, "Riga applicabile"):
            cs = K.uniq_cells(t.rows[1])
            A = q["sup_comp"]
            riga = ("A < 500 m²" if A < 500 else "500 ≤ A < 1000" if A < 1000 else "1000 ≤ A < 2500" if A < 2500 else
                    "2500 ≤ A < 5000" if A < 5000 else "5000 ≤ A < 10000" if A < 10000 else "A ≥ 10000")
            K.fill_tokens(cs[0], [IT(A, 0), riga]); K.fill_tokens(cs[1], [IT(q["dq1"])])
        for _, t in K.find_tables(b.doc, "Classe adottata e motivazione"):
            cs = K.uniq_cells(t.rows[1])
            K.fill_tokens(cs[0], [q["classe_dq2"], f"δα {r['dalfa']}, sostanze/lavorazioni pericolose: {'sì' if b.d('sostanze_significative') or b.d('lavorazioni_pericolose') else 'no'}"])
            K.fill_tokens(cs[1], [IT(q["dq2"])])
        for _, t in K.find_tables(b.doc, "Misure applicate"):
            cs = K.uniq_cells(t.rows[1])
            K.fill_tokens(cs[0], ["nessuna misura S.6/S.5/S.7/S.8/S.9 dei livelli indicati in Tab. S.2-8 (assunto cautelativo)"])
            K.fill_tokens(cs[1], [IT(q["sdn"])])
        for _, t in K.find_tables(b.doc, "Formula"):
            cs = K.uniq_cells(t.rows[1])
            K.fill_tokens(cs[1], [IT(q["dq1"]), IT(q["dq2"]), IT(q["sdn"]), IT(q["qf"], 0)])
            K.fill_tokens(cs[2], [IT(q["qfd"], 0)])
        for _, t in K.find_tables(b.doc, "CLASSE MINIMA DI RESISTENZA AL FUOCO ADOTTATA"):
            rei = q["rei"]
            K.fill_tokens(t.rows[0].cells[0], ["nessun requisito" if isinstance(rei, str) else f"{rei}"])

    # ---- S.1–S.10: criteri, livelli, cruscotto --------------------------------------------------------------
    tabelle = parse_criteri_tabelle(b.doc)
    from ..regole import expand_rvita, rvita_in
    for nmis, (tbl, liv0) in tabelle.items():
        dett = _dettaglio(stra.get(nmis), tbl)
        if nmis == 1:   # S.1: criteri descrittivi, valutati per Rvita
            vie, altri = ST.s1_livello(r["rvita"])
            for L in liv0:
                for ri, txt in L.condizioni:
                    if "Rvita =" in txt:
                        ins = expand_rvita(txt.split("Rvita =")[1])
                        dett[(L.nome, ri)] = rvita_in(r["rvita"], ins) if txt.startswith(("Vie", "Altri")) else None
                    elif "tipicamente" in txt:
                        dett[(L.nome, ri)] = r["rvita"].startswith("A")
                    else:
                        dett[(L.nome, ri)] = None
        for (nome, ri), v in dett.items():
            cell = K.uniq_cells(tbl.rows[ri])[2]
            K.set_cell(cell, "☒" if v is True else ("—" if v is False else "☐"))
    # paragrafi "Livello di prestazione ADOTTATO"
    k = 0
    for p in b.doc.paragraphs:
        if p.text.strip().startswith("Livello di prestazione ADOTTATO"):
            k += 1
            o = stra.get(k)
            if o:
                extra = f" ({o['dettaglio']})" if o.get("dettaglio") else ""
                K.fill_tokens(p, [f"{o['livello']}{extra}"])
    for _, t in K.find_tables(b.doc, "Misura"):
        if len(t.rows) == 11 and len(K.uniq_cells(t.rows[0])) == 4:
            for i in range(1, 11):
                cs = K.uniq_cells(t.rows[i])
                o = stra.get(i)
                if o:
                    K.set_cell(cs[2], o["livello"])
                    K.set_cell(cs[3], ("Da confermare in sopralluogo" if o["dubbi"] else "") or "—")
    b.checklist()
    # soluzione adottata: Conforme se tutta la check-list della misura è C/NA
    sol = [t for _, t in K.find_tables(b.doc, "Parametro") if len(t.rows) == 2 and K.ctext(t.rows[1].cells[0]).startswith("Soluzione adottata")]
    for i, t in enumerate(sol, start=1):
        if b.esito_misura(f"S.{i}") == "C":
            K.tick(t.rows[1].cells[1], "Conforme")
            K.fill_tokens(t.rows[1].cells[1], ["—"])

    # ---- S.4 scheda dimensionamento -------------------------------------------------------------------------
    esodo = {e["chiave"]: e for e in ST.verifiche_esodo(caso)}
    piani = int(b.d("piani_n", 1) or 1)
    for p in b.doc.paragraphs:
        t = p.text.strip()
        if t.startswith("☐ Esodo simultaneo"):
            K.tick(p, "Esodo simultaneo")
        elif t.startswith("N° uscite richieste"):
            K.fill_tokens(p, [str(dim["uscite_min"]), fmt_num(b.d("n_uscite"))])
            if "uscite" in esodo:
                K.tick(p, "Sì" if esodo["uscite"]["ok"] else "No")
        elif t.startswith("Corridoio cieco presente"):
            if b.d("lcc_misurata") not in (None, ""):
                K.tick(p, "Sì"); K.fill_tokens(p, [r["rvita"], fmt_num(dim["occupanti"]), fmt_num(b.d("lcc_misurata"))])
                K.tick(p, "OK" if esodo["lcc"]["ok"] else "NC")
            elif b.d("lcc_misurata") is None and b.d("n_uscite") is not None:
                K.tick(p, "No")
        elif t.startswith("Rvita adottato: ___"):
            K.fill_tokens(p, [r["rvita"], fmt_num(dim["les_max"]), fmt_num(b.d("les_misurata"))])
            if "les" in esodo:
                K.tick(p, "OK" if esodo["les"]["ok"] else "NC")
        elif t.startswith("LV calcolata"):
            lu = dim.get("lu_vert")
            if lu and piani > 1:
                K.fill_tokens(p, [IT(lu), fmt_num(dim["occupanti"]), fmt_num(dim.get("lv_calcolata_mm"))])
    for _, t in K.find_tables(b.doc, "Condizione"):
        if len(t.rows) == 4 and K.ctext(t.rows[0].cells[1]) == "Affollamento":
            occ = dim["occupanti"]
            sel = 1 if occ > 500 else (2 if (r["rvita"][0] == "B" and dim["densita"] > 0.4 and occ > 150) else 3)
            if "uscite" in esodo:
                K.choose(K.uniq_cells(t.rows[sel])[3], "OK" if esodo["uscite"]["ok"] else "NC")
    for _, t in K.find_tables(b.doc, "Rvita"):
        hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
        if len(hdr) == 6 and hdr[1].startswith("LU"):
            g = re.sub(r"^Ci{1,3}(\d)$", r"C\1", r["rvita"])
            for rw in t.rows[1:]:
                cs = K.uniq_cells(rw)
                if g in [x.strip() for x in K.ctext(cs[0]).split(",")] and dim.get("lu_oriz"):
                    K.fill_tokens(cs[2], [fmt_num(dim["occupanti"])]); K.fill_tokens(cs[3], [fmt_num(dim.get("lo_calcolata_mm"))])
                    K.fill_tokens(cs[4], [])
                    if b.d("largh_installata_mm"):
                        K.fill_tokens(cs[5], [fmt_num(b.d("largh_installata_mm"))])
                    K.shade(cs[0], K.OK_FILL)
        if len(hdr) == 8 and hdr[1] == "1 piano" and piani > 1 and dim.get("lu_vert"):
            g = {"C1": "B1", "E1": "B1", "C2": "B2", "D1": "B2", "E2": "B2", "C3": "B3", "D2": "B3", "E3": "B3"}.get(
                re.sub(r"^Ci{1,3}(\d)$", r"C\1", r["rvita"]), r["rvita"])
            for rw in t.rows[1:]:
                cs = K.uniq_cells(rw)
                if g in K.ctext(cs[0]).replace(",", " ").split():
                    K.fill_tokens(cs[7], [IT(dim["lu_vert"])]); K.shade(cs[0], K.OK_FILL)

    # ---- S.6 scheda presidi ----------------------------------------------------------------------------------
    ea = dim["estintori_A"]
    for _, t in K.find_tables(b.doc, "Rvita"):
        hdr = [K.ctext(x) for x in K.uniq_cells(t.rows[0])]
        if len(hdr) == 5 and hdr[1].startswith("Max distanza"):
            g = re.sub(r"^Ci{1,3}(\d)$", r"C\1", r["rvita"])
            for rw in t.rows[1:]:
                cs = K.uniq_cells(rw)
                if g in [x.strip() for x in K.ctext(cs[0]).split(",")]:
                    K.fill_tokens(cs[4], [f"n. ≥ {ea['n_min']} (sup. {dim_sup(b)} m² / area coperta r = {ea['dist_max']} m; ≥ 1 per piano)"])
                    K.shade(cs[0], K.OK_FILL)
                else:
                    K.set_cell(cs[4], "—")
    eb = dim["estintori_B"]
    for _, t in K.find_tables(b.doc, "Quantità liquido infiam. (L)"):
        lit = float(b.d("litri_infiammabili", 0) or 0)
        for rw in t.rows[1:]:
            cs = K.uniq_cells(rw)
            K.fill_tokens(cs[4], [fmt_num(lit) if lit else "N.A."])
    for _, t in K.find_tables(b.doc, "Piano/Compartimento"):
        cs = K.uniq_cells(t.rows[1])
        K.fill_tokens(cs[0], ["Intera attività"]); K.fill_tokens(cs[1], [str(ea["n_min"])])
        K.fill_tokens(cs[2], [fmt_num(b.d("n_estintori"))]); K.fill_tokens(cs[3], [str(eb["n"]) if eb else "N.A."])
        if "estintori" in esodo:
            K.choose(cs[4], "OK" if esodo["estintori"]["ok"] else "NC")
        for j in range(len(t.rows) - 1, 1, -1):
            K.delete_row(t, j)

    # ---- VRI-10 classificazione ------------------------------------------------------------------------------
    for _, t in K.find_tables(b.doc, "Classificazione"):
        if len(t.rows) == 2 and "BASSO" in K.ctext(t.rows[1].cells[0]):
            cs = K.uniq_cells(t.rows[1])
            K.choose(cs[0], "BASSO" if n["rischio_basso"] else "NON BASSO")
            K.fill_tokens(cs[1], ["; ".join(ko) if ko else "tutti i requisiti dell'Allegato I soddisfatti (VRI-2.1)"
                                  if n["rischio_basso"] else "almeno un requisito dell'Allegato I non soddisfatto o non verificabile (VRI-2.1)"])
    f = n["formazione"]
    b.par_fill("Livello di formazione applicabile", [f"Livello {f['livello']} ({f['ore']} ore{'; ' + f['agg'] if f['agg'] != '—' else ''})"])
    for _, t in K.find_tables(b.doc, "Classificazione"):
        if len(t.rows) == 4:
            row = {1: 1, 2: 2, 3: 3}[f["livello"]]
            K.shade(K.uniq_cells(t.rows[row])[0], K.OK_FILL); K.shade(K.uniq_cells(t.rows[row])[1], K.OK_FILL)
    b.azioni(); b.periodicita(); b.sottoscrizioni()
    b.label_fill("DATI DEL DOCUMENTO", {"Data di emissione": nz(b.d("data_emissione"), None) or __import__("datetime").date.today().strftime("%d/%m/%Y"),
                                       "Revisione n.": b.revisione_str()})
    pend = b.finalize()
    b.salva(out)
    return pend


def dim_sup(b):
    return fmt_num(b.d("superficie_mq")) or "—"


def _dettaglio(o, tbl):
    """Rivaluta le condizioni per colorare l'Esito della tabella dei criteri."""
    res = {}
    if not o:
        return res
    for k, v in (o.get("conds") or {}).items():
        a, bb = k.split("|")
        res[(a, int(bb))] = v
    return res
