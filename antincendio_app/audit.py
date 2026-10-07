"""Audit dei template: controlli di coerenza interna e refusi. Non sostituisce la verifica sul testo ufficiale della norma."""
from __future__ import annotations
import re
from pathlib import Path
import docx
from . import docxkit as K

TPL = Path(__file__).resolve().parents[1] / "templates"
PATTERN = {r"§\s*\d": "riferimento «§» non allineato alla numerazione VRI-x", r"\(\)": "riferimento vuoto «()»",
           r"Idirizzo": "refuso «Idirizzo»", r"ITERVENTO": "refuso «ITERVENTO»", r"\bpiú\b": "accento errato «piú»", r"Cap\. 5\b": "riferimento «VRI Cap. 5» inesistente"}


def testi(d):
    t = [p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in K.uniq_cells(r)]
    for sec in d.sections:
        t += [p.text for p in sec.header.paragraphs + sec.footer.paragraphs]
    return t


def verifica() -> list[str]:
    out = []
    docs = {f.name: docx.Document(str(f)) for f in sorted(TPL.glob("*.docx"))}
    for nome, d in docs.items():
        tx = testi(d)
        for rx, desc in PATTERN.items():
            h = [t for t in tx if re.search(rx, t)]
            if h:
                out.append(f"{nome}: {desc} ({len(h)} occorrenze) — es. «{h[0][:90].strip()}»")
    # coerenza tra Minicodice e Codice sui requisiti di basso rischio
    mini = "\n".join(testi(docs["VRI_Minicodice_template_2.docx"])) if "VRI_Minicodice_template_2.docx" in docs else ""
    cod = "\n".join(testi(docs["VRI_Codice_Integrale_RTO_completa_2.docx"])) if "VRI_Codice_Integrale_RTO_completa_2.docx" in docs else ""
    if "C.f" in cod and "C.f" not in mini:
        out.append("Minicodice vs Codice: il Codice verifica 6 requisiti di basso rischio (C.a–C.f: occupanti, superficie, quote, qf≤900, sostanze, lavorazioni) "
                   "mentre la VRI-2.3 del Minicodice ne dimostra solo 3 numerici + 2 di quadro; qf/sostanze/lavorazioni sono trattati altrove (VRI-5/6). Allineare la VRI-2.3.")
    return out
