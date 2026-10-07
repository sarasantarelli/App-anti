"""Agente NORME: corpus normativo/tecnico fornito dall'utente (norme, libri di valutazione, dispense), indicizzato in locale.

- Si aggiunge dall'interfaccia (scheda «Norme e libri») o copiando file in `norme/` (PDF, Word, TXT, MD).
- L'indice è salvato su disco (`norme/.indice/`): un libro di centinaia di pagine si indicizza una volta sola.
- PDF scansionati: OCR locale pagina per pagina (più lento; limite OCR_MAX_PAGINE per file).
- Uso: ricerca per parole chiave con citazione (file, pagina) e allegato dei passi pertinenti ai rilievi della Relazione di controllo.
L'app non inventa articoli: cita solo ciò che è stato caricato.
"""
from __future__ import annotations
import hashlib, json, math, os, re, threading
from collections import Counter
from pathlib import Path

NORME_DIR = Path(os.environ.get("NORME_DIR", Path(__file__).resolve().parents[2] / "norme"))
IDX_DIR = NORME_DIR / ".indice"
EXT = (".pdf", ".docx", ".txt", ".md")
OCR_MAX_PAGINE = int(os.environ.get("OCR_MAX_PAGINE", "400"))
STOP = set("il lo la i gli le un una di del della dei delle da in con su per tra fra e ed o che si non è sono al allo alla ai agli alle dal dalla nel nella nei nelle sul sulla".split())
_lock = threading.Lock()
_stato = {"in_corso": [], "errori": {}}
_cache = {"sig": None, "docs": None, "df": None}


def _tok(t: str) -> list[str]:
    return [w for w in re.findall(r"[a-zàèéìòù0-9]{3,}", t.lower()) if w not in STOP]


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ ()-]", "_", Path(name).name)[:150]


def file_corpus() -> list[Path]:
    NORME_DIR.mkdir(parents=True, exist_ok=True)
    return sorted(f for f in NORME_DIR.glob("*") if f.is_file() and f.suffix.lower() in EXT and not f.name.upper().startswith("LEGGIMI"))


def _idx_path(f: Path) -> Path:
    return IDX_DIR / (hashlib.sha1(f.name.encode()).hexdigest()[:16] + ".json")


def _pagine(f: Path):
    ext = f.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        from .ingestion import ocr_image
        import tempfile
        with pdfplumber.open(f) as pdf:
            for i, pg in enumerate(pdf.pages, start=1):
                t = pg.extract_text() or ""
                if len(t.strip()) < 30 and i <= OCR_MAX_PAGINE:
                    try:
                        with tempfile.TemporaryDirectory() as td:
                            img = Path(td) / "p.png"
                            pg.to_image(resolution=130).save(img)
                            t = ocr_image(str(img))
                    except Exception:
                        pass
                yield i, t
    elif ext == ".docx":
        import docx
        d = docx.Document(str(f))
        righe = [p.text for p in d.paragraphs] + [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
        for k in range(0, len(righe), 40):
            yield k // 40 + 1, "\n".join(righe[k:k + 40])
    else:
        t = f.read_text(errors="ignore")
        rr = t.splitlines()
        for k in range(0, len(rr), 60):
            yield k // 60 + 1, "\n".join(rr[k:k + 60])


def _indicizza_file(f: Path):
    sig = {"size": f.stat().st_size, "mtime": int(f.stat().st_mtime)}
    ip = _idx_path(f)
    if ip.exists():
        try:
            j = json.loads(ip.read_text(encoding="utf-8"))
            if j.get("sig") == sig:
                return j
        except Exception:
            pass
    pas = []
    for pg, testo in _pagine(f):
        righe = [r.strip() for r in testo.splitlines() if r.strip()]
        for k in range(0, len(righe), 8):
            blocco = " ".join(righe[k:k + 10])
            if len(blocco) > 60:
                pas.append({"p": pg, "t": blocco})
    j = {"file": f.name, "sig": sig, "pagine": (pas[-1]["p"] if pas else 0), "passaggi": pas}
    IDX_DIR.mkdir(parents=True, exist_ok=True)
    ip.write_text(json.dumps(j, ensure_ascii=False), encoding="utf-8")
    return j


def indicizza_tutto():
    for f in file_corpus():
        with _lock:
            _stato["in_corso"].append(f.name)
        try:
            _indicizza_file(f)
            _stato["errori"].pop(f.name, None)
        except Exception as e:
            _stato["errori"][f.name] = str(e)[:200]
        finally:
            with _lock:
                if f.name in _stato["in_corso"]:
                    _stato["in_corso"].remove(f.name)


def indicizza_async():
    if _stato["in_corso"]:
        return
    threading.Thread(target=indicizza_tutto, daemon=True).start()


def stato() -> dict:
    return {"in_corso": list(_stato["in_corso"]), "errori": dict(_stato["errori"])}


def elenco() -> list[dict]:
    out = []
    for f in file_corpus():
        ip, j = _idx_path(f), None
        if ip.exists():
            try:
                j = json.loads(ip.read_text(encoding="utf-8"))
            except Exception:
                pass
        out.append({"file": f.name, "mb": round(f.stat().st_size / 1e6, 2), "indicizzato": bool(j and j.get("sig", {}).get("size") == f.stat().st_size),
                    "pagine": (j or {}).get("pagine"), "passaggi": len((j or {}).get("passaggi", [])),
                    "in_corso": f.name in _stato["in_corso"], "errore": _stato["errori"].get(f.name)})
    return out


def aggiungi(nome: str, contenuto: bytes) -> str:
    nome = _safe(nome)
    if Path(nome).suffix.lower() not in EXT:
        raise ValueError("formato non supportato (PDF, DOCX, TXT, MD)")
    NORME_DIR.mkdir(parents=True, exist_ok=True)
    (NORME_DIR / nome).write_bytes(contenuto)
    return nome


def rimuovi(nome: str):
    f = NORME_DIR / _safe(nome)
    if f.exists() and f.suffix.lower() in EXT:
        f.unlink()
    ip = _idx_path(f)
    if ip.exists():
        ip.unlink()


def _carica():
    fl = file_corpus()
    sig = tuple((f.name, f.stat().st_size, int(f.stat().st_mtime)) for f in fl)
    if _cache["sig"] == sig and _cache["docs"] is not None:
        return _cache["docs"], _cache["df"]
    docs = []
    for f in fl:
        ip = _idx_path(f)
        if not ip.exists():
            continue                 # non ancora indicizzato: lo farà indicizza_async()
        try:
            j = json.loads(ip.read_text(encoding="utf-8"))
        except Exception:
            continue
        for p in j["passaggi"]:
            docs.append((f.name, p["p"], p["t"], Counter(_tok(p["t"]))))
    df = Counter()
    for *_, c in docs:
        df.update(c.keys())
    _cache.update(sig=sig, docs=docs, df=df)
    return docs, df


def cerca(query: str, n: int = 3) -> list[dict]:
    docs, df = _carica()
    q = _tok(query)
    N = len(docs)
    if not docs or not q:
        return []
    frase = query.lower().strip()
    sc = []
    for f, pg, testo, c in docs:
        s = sum((1 + math.log(c[w])) * math.log(1 + N / (1 + df[w])) for w in q if c.get(w))
        if s > 0:
            if len(frase) > 6 and frase in testo.lower():
                s *= 1.5
            sc.append((s, f, pg, testo))
    sc.sort(reverse=True)
    return [{"file": f, "pagina": pg, "testo": t[:520], "punteggio": round(s, 2)} for s, f, pg, t in sc[:n]]


def arricchisci(caso) -> None:
    """Allega ai rilievi con riferimento normativo il passo più pertinente del corpus caricato."""
    docs, _ = _carica()
    if not docs:
        return
    for f in caso.findings:
        if f.riferimento and "Riscontro nel corpus" not in f.messaggio:
            r = cerca(f.riferimento + " " + f.messaggio[:120], 1)
            if r and r[0]["punteggio"] > 6:
                f.messaggio += f" [Riscontro nel corpus: {r[0]['file']} p.{r[0]['pagina']}: «{r[0]['testo'][:220]}…»]"
