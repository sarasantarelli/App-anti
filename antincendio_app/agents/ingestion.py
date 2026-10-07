"""Agente LETTURA: estrae testo da documenti, foto e video. Tutto in locale, nessuna chiave API.

- PDF: testo nativo (pdfplumber); se la pagina è scansione -> OCR locale (RapidOCR/ONNX).
- DOCX/XLSX/TXT/CSV: lettura diretta.
- Immagini: OCR locale (cartelli, targhe estintori, planimetrie con quote scritte).
- Video: fotogrammi campionati con ffmpeg -> OCR -> deduplica.
Limite dichiarato: senza modello di visione NON si riconoscono oggetti (es. estintore non etichettato);
si leggono solo testi/cartelli e si registra il fotogramma come evidenza da verificare dal tecnico.
"""
from __future__ import annotations
import subprocess, tempfile, shutil
from pathlib import Path

_OCR = None


def _ocr_engine():
    global _OCR
    if _OCR is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _OCR = RapidOCR()
        except Exception:  # OCR non disponibile: degradazione dichiarata
            _OCR = False
    return _OCR


def ocr_image(path: str) -> str:
    eng = _ocr_engine()
    if not eng:
        return ""
    res, _ = eng(str(path))
    return "\n".join(r[1] for r in (res or []))


def _pdf(path: Path) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            t = page.extract_text() or ""
            if len(t.strip()) < 30:  # probabile scansione
                with tempfile.TemporaryDirectory() as td:
                    img = Path(td) / "p.png"
                    page.to_image(resolution=150).save(img)
                    t = ocr_image(str(img))
            out.append(t)
    return "\n".join(out)


def _docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb:
        parts.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                parts.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(parts)


def _video(path: Path, every_s: int = 3, max_frames: int = 40) -> str:
    from ..ambiente import trova_ffmpeg
    ff = trova_ffmpeg()
    if not ff:
        return ""
    texts, seen = [], set()
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([ff, "-loglevel", "error", "-i", str(path), "-vf",
                        f"fps=1/{every_s},scale=1280:-1", "-frames:v", str(max_frames),
                        f"{td}/f_%03d.jpg"], check=False)
        for f in sorted(Path(td).glob("f_*.jpg")):
            for line in ocr_image(str(f)).splitlines():
                k = line.strip().lower()
                if k and k not in seen:
                    seen.add(k)
                    texts.append(line.strip())
    return "\n".join(texts)


READERS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx,
           ".txt": lambda p: p.read_text(errors="ignore"), ".md": lambda p: p.read_text(errors="ignore"),
           ".csv": lambda p: p.read_text(errors="ignore")}
IMG = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
VID = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"}


def leggi_file(path: str | Path) -> dict:
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext in READERS:
            tipo, testo = "documento", READERS[ext](p)
        elif ext in IMG:
            tipo, testo = "foto", ocr_image(str(p))
        elif ext in VID:
            tipo, testo = "video", _video(p)
        else:
            return {"file": p.name, "tipo": "non supportato", "testo": "", "nota": f"formato {ext} non supportato"}
    except Exception as e:
        return {"file": p.name, "tipo": "errore", "testo": "", "nota": f"lettura fallita: {e}"}
    nota = "" if testo.strip() else "nessun testo leggibile: esaminare manualmente"
    return {"file": p.name, "tipo": tipo, "testo": testo, "nota": nota}
