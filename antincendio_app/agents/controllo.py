"""Agente CONTROLLO: verifica completezza, coerenza e incrocio con la norma. Non corregge in silenzio:
ogni anomalia diventa un rilievo (Finding) e, se risolvibile solo dal tecnico, una domanda."""
from __future__ import annotations
from ..models import Caso
from .. import regole as R

AGENTE = "Controllo"

OBBLIGATORI = [
    ("ragione_sociale", "Ragione sociale"), ("indirizzo", "Indirizzo della sede/unità produttiva"),
    ("datore_lavoro", "Datore di lavoro"), ("rspp", "RSPP"), ("attivita_descrizione", "Descrizione reale dell'attività"),
    ("tipologia_attivita", "Tipologia di attività"), ("superficie_mq", "Superficie lorda complessiva"),
    ("quota_min", "Quota del piano più basso"), ("quota_max", "Quota del piano più alto"),
    ("occupanti_lavoratori", "Numero di lavoratori"),
]


def domande(c: Caso) -> list[dict]:
    q = []
    for k, lab in OBBLIGATORI:
        if c.get(k) in (None, "", []):
            q.append({"chiave": k, "domanda": f"Manca: {lab}", "gravita": "bloccante" if k in ("superficie_mq", "occupanti_lavoratori", "tipologia_attivita") else "completamento"})
    if c.get("sostanze_significative") is None:
        q.append({"chiave": "sostanze_significative", "domanda": "Sono presenti sostanze/miscele pericolose in quantità significative? (decide il percorso Minicodice/Codice)", "gravita": "bloccante"})
    if c.get("lavorazioni_pericolose") is None:
        q.append({"chiave": "lavorazioni_pericolose", "domanda": "Si effettuano lavorazioni pericolose ai fini dell'incendio? (decide il percorso)", "gravita": "bloccante"})
    qf = c.esito.get("rischio", {}).get("qf")
    if qf is None:
        q.append({"chiave": "qf_mj_m2", "domanda": "Carico d'incendio qf non determinabile: inserire qf [MJ/m²] o l'elenco dei materiali con quantità", "gravita": "bloccante"})
    return q


def pre(c: Caso) -> Caso:
    c.esito["domande"] = domande(c)
    return c


def post(c: Caso, voci, risposte) -> Caso:
    n, r, dim = c.esito["normativo"], c.esito["rischio"], c.esito["dimensionamento"]
    sup = float(c.get("superficie_mq", 0) or 0)
    occ = n["occupanti"]
    # 1. coerenza locali / superficie
    loc = c.get("locali") or []
    if loc:
        tot = sum(float(l.get("mq") or 0) for l in loc)
        if sup and abs(tot - sup) > 0.05 * sup:
            c.add(AGENTE, "attenzione", f"La somma delle superfici dei locali ({tot:g} m²) differisce dalla superficie complessiva dichiarata ({sup:g} m²).", "VRI-4.1")
    # 2. densità di affollamento anomala
    if sup and occ / sup > 1.0:
        c.add(AGENTE, "attenzione", f"Densità di affollamento {occ / sup:.2f} p/m² molto elevata: verificare il conteggio degli occupanti.", "Codice S.4 Tab. S.4-30")
    # 3. soglie prossime ai limiti (instabilità della classificazione)
    if n["ramo"] != "MINICODICE" and occ and sup:
        near = []
        if 90 <= occ <= 110: near.append(f"affollamento {occ:g} vicino a 100")
        if 900 <= sup <= 1100: near.append(f"superficie {sup:g} m² vicina a 1.000")
        if near:
            c.add(AGENTE, "info", "Classificazione sensibile alle soglie: " + "; ".join(near) + ". Documentare con rilievo e conteggio accurati.", "DM 3/9/2021 All. I")
    # 4. rischio basso con elementi speciali
    if n["rischio_basso"] and r["qf"] and r["qf"] > 900:
        c.add(AGENTE, "critico", "Minicodice con qf > 900 MJ/m²: incoerente con il requisito di basso rischio.", "DM 3/9/2021 All. I")
    # 5. presidi richiesti dalla strategia ma non citati nei documenti
    st = c.esito.get("strategia", {})
    pres = set(c.get("presidi_rilevati", []) or [])
    if st.get("modo") == "codice":
        m = st["misure"]
        if m.get(7, {}).get("livello") in ("II", "III", "IV") and "rivelazione/allarme incendio" not in pres:
            c.add(AGENTE, "attenzione", f"S.7 livello {m[7]['livello']}: richiede IRAI, non citato nei documenti. Verificare in sopralluogo o pianificare l'impianto.", "Codice S.7")
        if m.get(6, {}).get("livello") in ("III", "IV") and "idranti/naspi" not in pres:
            c.add(AGENTE, "attenzione", f"S.6 livello {m[6]['livello']}: richiesta rete idranti, non citata nei documenti.", "Codice S.6")
    if "estintori" not in pres:
        c.add(AGENTE, "info", "Nei documenti non risultano estintori: verificare in sopralluogo (dotazione minima sempre richiesta).", "Allegato I DM 3/9/2021 M.4 / Codice S.6")
    # 6. formazione
    f = n["formazione"]
    c.add(AGENTE, "info", f"Formazione addetti: Livello {f['livello']} ({f['ore']} ore) in funzione della classificazione ({'basso' if n['rischio_basso'] else 'non basso'}).", "DM 2/9/2021; Accordo S-R 17/4/2025")
    # 7. confini di responsabilità
    if n["soggetta_dpr151"]:
        c.add(AGENTE, "info", "Attività soggetta ai controlli VVF: la VRI è in raccordo con la pratica; progetto, SCIA e asseverazioni restano del professionista antincendio iscritto.", "DPR 151/2011")
    # 8. completezza rilievi in sito
    tot = len(voci)
    risp = sum(1 for v in voci if risposte.get(v.id, {}).get("esito"))
    c.esito["completezza"] = {"voci": tot, "risposte": risp, "percento": round(100 * risp / tot) if tot else 100}
    if risp < tot:
        c.add(AGENTE, "attenzione", f"Sopralluogo incompleto: {tot - risp} voci su {tot} senza esito. Il documento resta in BOZZA finché non sono completate.", "VRI — checklist")
    crit = [f for f in c.findings if f.livello == "critico"]
    bloc = [d for d in c.esito.get("domande", []) if d["gravita"] == "bloccante"]
    c.esito["stato"] = "PRONTO PER LA FIRMA" if (risp == tot and not crit and not bloc) else "BOZZA"
    return c


def verifica_documento(path, c: Caso) -> list[str]:
    """Rilegge il .docx generato e verifica che i dati chiave siano coerenti con l'analisi."""
    import docx
    d = docx.Document(str(path))
    txt = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                txt += "\n" + cell.text
    problemi = []
    n, r = c.esito["normativo"], c.esito["rischio"]
    if c.get("ragione_sociale") and str(c.get("ragione_sociale")).lower() not in txt.lower():
        problemi.append("ragione sociale non presente nel documento")
    if n["template"] == "codice" and "VRI_" in str(path) and r["rvita"] not in txt:
        problemi.append(f"Rvita {r['rvita']} non riportato")
    if "GUIDA ALLA COMPILAZIONE" in txt:
        problemi.append("restano box di guida alla compilazione")
    return problemi
