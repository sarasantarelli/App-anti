"""Utility per compilare i template .docx ufficiali PRESERVANDO font, colori e layout.

Principio: non si ricrea mai il documento; si apre il template dell'utente e si scrive
nelle sue celle/paragrafi riusando le proprietà di formattazione delle run esistenti.
"""
from __future__ import annotations
import copy, re
from docx.document import Document as _Doc
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BOX, TICK = "☐", "☒"
PENDING_FILL = "FAE6C0"   # palette template: avvertenza -> campo da completare
NC_FILL = "F2DEDE"
OK_FILL = "DFF0D8"
PLACEHOLDER = re.compile(r"\[(?!\d+\]|m²\]|MJ/m²\])[^\]]{1,160}\]|_{3,}")


def blocks(doc):
    """Paragrafi e tabelle nell'ordine del documento."""
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            yield Paragraph(el, doc)
        elif el.tag == qn("w:tbl"):
            yield Table(el, doc)


def uniq_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if c._tc in seen:
            continue
        seen.add(c._tc)
        out.append(c)
    return out


def ctext(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def _runs(cell):
    return [r for p in cell.paragraphs for r in p.runs]


def _pick_run(p):
    """Run di riferimento per la formattazione: l'ultima con testo, altrimenti l'ultima."""
    base = None
    for r in p.runs:
        if r.text:
            base = r
    if base is None and p.runs:
        base = p.runs[-1]
    return base


def _write_par(p, text):
    base = _pick_run(p)
    if base is None:
        base = p.add_run("")
    keep_el = base._element
    for r in list(p.runs):
        if r._element is not keep_el:
            r._element.getparent().remove(r._element)
    base.text = text
    return base


def set_cell(cell: _Cell, text: str, fill: str | None = None, bold: bool | None = None):
    """Scrive `text` nella cella riusando la formattazione dell'ultima run con testo."""
    paras = cell.paragraphs
    # formattazione di riferimento: ultima run con testo in qualunque paragrafo
    ref = None
    for p in paras:
        r = _pick_run(p)
        if r is not None and r.text:
            ref = r
    p0 = paras[0]
    if ref is not None and ref._element.getparent() is not p0._element:
        # porta la formattazione di riferimento nel primo paragrafo
        import copy as _c
        newr = _c.deepcopy(ref._element)
        p0._element.append(newr)
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    base = _write_par(p0, text)
    if bold is not None:
        base.font.bold = bold
    if fill:
        shade(cell, fill)


def shade(cell: _Cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn("w:shd")):
        tcPr.remove(s)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_par(p: Paragraph, text: str):
    _write_par(p, text)


def tick(container, label: str, state: bool = True, nth: int = 0) -> bool:
    """Marca '☐ label' -> '☒ label' (prima occorrenza) in una cella o paragrafo."""
    paras = container.paragraphs if hasattr(container, "paragraphs") else [container]
    pat = re.compile(r"☐(\s*)(" + re.escape(label) + r")(?![A-Za-z])")
    seen = 0
    for p in paras:
        full = p.text
        ms = list(pat.finditer(full))
        if len(ms) > nth - seen if False else False:
            pass
        if not ms:
            continue
        if nth - seen >= len(ms):
            seen += len(ms)
            continue
        k = nth - seen
        # prova modifica in-place su run singola
        idx = 0
        for r in p.runs:
            rms = list(pat.finditer(r.text))
            if idx + len(rms) > k:
                target = rms[k - idx]
                r.text = r.text[:target.start()] + (TICK if state else BOX) + target.group(1) + target.group(2) + r.text[target.end():]
                return True
            idx += len(rms)
        # etichetta spezzata su più run: riscrive il paragrafo
        m = ms[k]
        set_par(p, full[:m.start()] + (TICK if state else BOX) + m.group(1) + m.group(2) + full[m.end():])
        return True
    return False


def untick_all(container):
    for p in (container.paragraphs if hasattr(container, "paragraphs") else [container]):
        for r in p.runs:
            if TICK in r.text:
                r.text = r.text.replace(TICK, BOX)


def choose(container, label: str):
    """Esclusiva: azzera le altre scelte e marca `label`."""
    untick_all(container)
    return tick(container, label)


def fill_tokens(container, values: list[str]) -> int:
    """Sostituisce in ordine i segnaposto [..] / ___ con `values` (None = lascia)."""
    paras = container.paragraphs if hasattr(container, "paragraphs") else [container]
    it, n = iter(values), 0
    for p in paras:
        for r in p.runs:
            def rep(m):
                nonlocal n
                try:
                    v = next(it)
                except StopIteration:
                    return m.group(0)
                if v is None:
                    return m.group(0)
                n += 1
                return str(v)
            if PLACEHOLDER.search(r.text):
                r.text = PLACEHOLDER.sub(rep, r.text)
    return n


def has_pending(container) -> bool:
    paras = container.paragraphs if hasattr(container, "paragraphs") else [container]
    return any(PLACEHOLDER.search(p.text) for p in paras)


def _dark(cell) -> bool:
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return False
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return False
    f = (shd.get(qn("w:fill")) or "FFFFFF").lstrip("#")
    try:
        r, g, b = int(f[0:2], 16), int(f[2:4], 16), int(f[4:6], 16)
    except ValueError:
        return False
    return (0.299 * r + 0.587 * g + 0.114 * b) < 140


def mark_pending(doc, skip_headers=True) -> int:
    """Evidenzia (palette COR) le celle che contengono ancora segnaposto: campi da completare."""
    n = 0
    for t in doc.tables:
        for row in t.rows:
            for c in uniq_cells(row):
                if has_pending(c):
                    if _dark(c):
                        for p in c.paragraphs:
                            for r in p.runs:
                                if PLACEHOLDER.search(r.text):
                                    r.font.highlight_color = 7
                                    r.font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)
                    else:
                        shade(c, PENDING_FILL)
                    n += 1
    for p in doc.paragraphs:
        if PLACEHOLDER.search(p.text):
            for r in p.runs:
                if PLACEHOLDER.search(r.text):
                    r.font.highlight_color = 7  # YELLOW
            n += 1
    return n


def clone_row(table: Table, idx: int):
    tr = table.rows[idx]._tr
    new = copy.deepcopy(tr)
    tr.addnext(new)
    return table.rows[idx + 1]


def ensure_rows(table: Table, template_row: int, total: int):
    """Garantisce `total` righe dati clonando la riga modello (le righe dati iniziano da template_row)."""
    have = len(table.rows) - template_row
    last = len(table.rows) - 1
    for _ in range(max(0, total - have)):
        clone_row(table, last)
        last += 1
    return table


def delete_row(table: Table, idx: int):
    tr = table.rows[idx]._tr
    tr.getparent().remove(tr)


def remove_block(b):
    el = b._element if hasattr(b, "_element") else b._tbl
    el.getparent().remove(el)


def strip_guides(doc, keep=False) -> int:
    """Rimuove i box 'GUIDA ALLA COMPILAZIONE' e le istruzioni al redattore: il documento esibito non le contiene."""
    if keep:
        return 0
    n = 0
    for b in list(blocks(doc)):
        if isinstance(b, Table):
            if len(b.rows) == 1 and len(b.columns) == 1:
                tx0 = ctext(b.rows[0].cells[0])
                if tx0.upper().startswith("GUIDA ALLA COMPILAZIONE"):
                    remove_block(b); n += 1
                elif tx0.startswith("Duplicare questa scheda"):
                    remove_block(b); n += 1
        else:
            t = b.text.strip()
            if re.match(r"^(Duplicare|Compilare solo se|Compilare UNA|Compilare in ordine|✎|Compilare le ubicazioni|Compilare una riga)", t):
                remove_block(b); n += 1
    return n


def find_tables(doc, first_cell: str, startswith=True):
    out = []
    for i, t in enumerate(doc.tables):
        try:
            tx = ctext(t.rows[0].cells[0])
        except Exception:
            continue
        if (tx.startswith(first_cell) if startswith else first_cell in tx):
            out.append((i, t))
    return out


def label_map(table: Table):
    """Tabelle etichetta|valore: {etichetta: cella_valore}."""
    m = {}
    for row in table.rows:
        cs = uniq_cells(row)
        if len(cs) >= 2:
            m[ctext(cs[0])] = cs[1]
    return m


def heading_before(doc, table: Table) -> str:
    """Testo dell'ultimo paragrafo non vuoto che precede la tabella."""
    last = ""
    for b in blocks(doc):
        if isinstance(b, Table):
            if b._tbl is table._tbl:
                return last
        elif b.text.strip():
            last = b.text.strip()
    return last
