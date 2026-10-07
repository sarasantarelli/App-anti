import docx, pytest
from antincendio_app.coordinatore import esegui_tutto
from antincendio_app.agents import estrazione, ingestion
from antincendio_app.models import Caso
from antincendio_app import checklist as CL
from pathlib import Path

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

BASE = dict(ragione_sociale="Trattoria da Mario S.r.l.", indirizzo="Via Roma 12 — 00100 Roma (RM)", datore_lavoro="Mario Rossi",
            rspp="Sara Santarelli", rls="Anna Bianchi", superficie_mq=350, quota_min=0, quota_max=4, occupanti_lavoratori=6,
            occupanti_terzi=30, tipologia_attivita=["ristorazione"], aperta_pubblico=True, sostanze_significative=False,
            lavorazioni_pericolose=False, qf_mj_m2=450, piani_n=1)


def leggi_testo(path):
    d = docx.Document(str(path))
    t = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                t += "\n" + c.text
    return t


def run(tmp_path, dati, **kw):
    c, res = esegui_tutto(dati, kw.get("files", []), tmp_path, pdf=False, risposte=kw.get("risposte"))
    return c, res


def test_minicodice(tmp_path):
    c, res = run(tmp_path, BASE)
    assert c.esito["normativo"]["ramo"] == "MINICODICE" and c.esito["normativo"]["template"] == "minicodice"
    t = leggi_testo(tmp_path / res["documenti"][0]["file"])
    assert "TRATTORIA DA MARIO S.R.L." in t and "GUIDA ALLA COMPILAZIONE" not in t
    assert "BASSO RISCHIO" in t and res["stato"] == "BOZZA"
    assert "Rvita" in t   # il template menziona il profilo, ma non lo calcola
    assert all(not d["verifica"] for d in res["documenti"])


def test_codice_e_qfd(tmp_path):
    dati = dict(BASE, tipologia_attivita=["officina/produzione"], sostanze_significative=True, superficie_mq=900, occupanti_terzi=0,
                aperta_pubblico=False, sostanze=["liquidi infiammabili"], qf_mj_m2=1100, n_uscite=1, les_misurata=80)
    c, res = run(tmp_path, dati)
    assert c.esito["normativo"]["template"] == "codice"
    q = c.esito["rischio"]["qfd"]
    assert round(q["qfd"]) == round(q["dq1"] * q["dq2"] * q["sdn"] * 1100)
    t = leggi_testo(tmp_path / res["documenti"][0]["file"])
    assert f"{c.esito['rischio']['rvita']} ◄ adottato" in t
    # verifiche dimensionali: 1 uscita su >50 occupanti? (6 occ => 1 ok) ; Les 80 m > massimo => NC e azione
    assert any("lunghezza" in a["azione"].lower() for a in c.esito["azioni"])


def test_raccordo_e_scadenza(tmp_path):
    dati = dict(BASE, tipologia_attivita=["commercio"], superficie_mq=650, occupanti_terzi=60, scadenza_rinnovo="10/03/2020",
                categoria_dpr151="B", tipo_procedimento="SCIA art. 4")
    c, res = run(tmp_path, dati)
    assert c.esito["normativo"]["template"] == "raccordo" and c.esito["normativo"]["soggetta_dpr151"]
    assert any(a["priorita"] == 1 and "scaduto" in a["azione"].lower() for a in c.esito["azioni"])
    t = leggi_testo(tmp_path / res["documenti"][0]["file"])
    assert "in ritardo di" in t


def test_nc_diventa_azione_nel_documento(tmp_path):
    c0, _ = run(tmp_path / "a", BASE)
    voci = CL.estrai(docx.Document(str(Path(__file__).resolve().parents[1] / "templates" / "VRI_Minicodice_template_2.docx")), "minicodice")
    v = voci[8]
    c, res = run(tmp_path / "b", BASE, risposte={v.id: {"esito": "NC", "nota": "Uscita ostruita", "azione": "Liberare l'uscita XYZ"}})
    assert any("Liberare l'uscita XYZ" in a["azione"] for a in c.esito["azioni"])
    t = leggi_testo(tmp_path / "b" / res["documenti"][0]["file"])
    assert "Liberare l'uscita XYZ" in t and "Uscita ostruita" in t


def test_non_dichiara_nessuna_nc_se_sopralluogo_incompleto(tmp_path):
    c, res = run(tmp_path, BASE)
    t = leggi_testo(tmp_path / res["documenti"][0]["file"])
    assert "Nessuna non conformità rilevata" not in t


def test_lettura_e_estrazione_da_file(tmp_path):
    c, res = run(tmp_path, {}, files=[SAMPLES / "DVR_esempio.docx", SAMPLES / "foto_cartelli.png", SAMPLES / "video_sopralluogo.mp4"])
    assert c.get("ragione_sociale") == "Falegnameria Bianchi S.n.c."
    assert c.get("occupanti_lavoratori") == 12 and c.get("superficie_mq") == 780 and c.get("qf_mj_m2") == 1100
    assert "estintori" in c.get("presidi_rilevati") and "segnaletica sicurezza" in c.get("presidi_rilevati")
    assert c.dati["ragione_sociale"].fonte == "DVR_esempio.docx" and not c.dati["ragione_sociale"].confermato
    assert c.esito["stato"] == "BOZZA" and any(d["gravita"] == "bloccante" for d in c.esito["domande"])


def test_tutti_i_template_hanno_voci_stabili():
    base = Path(__file__).resolve().parents[1] / "templates"
    for f, k in (("VRI_Minicodice_template_2.docx", "minicodice"), ("VRI_Codice_Integrale_RTO_completa_2.docx", "codice"), ("VRI_Raccordo_CPI_template.docx", "raccordo")):
        a = CL.estrai(docx.Document(str(base / f)), k)
        b = CL.estrai(docx.Document(str(base / f)), k)
        assert len(a) > 20 and [x.id for x in a] == [x.id for x in b] and len({x.id for x in a}) == len(a)
