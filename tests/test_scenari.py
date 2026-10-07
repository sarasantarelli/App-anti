import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "scripts"))
import sweep


def test_giro_di_prove_senza_anomalie():
    assert sweep.main(8, 3) == 0


def test_ristorante_con_eventi_e_discriminanti(tmp_path):
    from antincendio_app.coordinatore import esegui_tutto
    base = dict(ragione_sociale="Locale X", indirizzo="Via 1", datore_lavoro="A", rspp="B", superficie_mq=150, quota_min=0, quota_max=3, occupanti_lavoratori=4,
                occupanti_terzi=40, tipologia_attivita=["ristorazione"], aperta_pubblico=True, sostanze_significative=False, lavorazioni_pericolose=False, qf_mj_m2=300)
    c, _ = esegui_tutto(dict(base, eventi_intrattenimento=True), [], tmp_path / "a", pdf=False)
    assert c.esito["normativo"]["soggetta_dpr151"] is False        # musica senza discriminanti: esclusa (V.15.1 c.2 lett. b)
    c, _ = esegui_tutto(dict(base, eventi_intrattenimento=True, pista_ballo=True), [], tmp_path / "b", pdf=False)
    assert c.esito["normativo"]["soggetta_dpr151"] is True          # pista da ballo: n. 65
    c, _ = esegui_tutto(dict(base, eventi_intrattenimento=True, pista_ballo=True, eventi_solo_aperto=True), [], tmp_path / "c", pdf=False)
    assert c.esito["normativo"]["soggetta_dpr151"] is False         # manifestazioni temporanee all'aperto: escluse


def test_piu_destinazioni_governa_il_profilo_piu_elevato(tmp_path):
    from antincendio_app.coordinatore import esegui_tutto
    d = dict(ragione_sociale="Hotel Y", indirizzo="Via 2", datore_lavoro="A", rspp="B", superficie_mq=500, quota_min=0, quota_max=8, occupanti_lavoratori=6,
             occupanti_terzi=20, tipologia_attivita=["ufficio"], aperta_pubblico=True, sostanze_significative=False, lavorazioni_pericolose=False, qf_mj_m2=300,
             locali=[{"nome": "Uffici", "mq": 200, "tipo": "ufficio"}, {"nome": "Camere", "mq": 300, "tipo": "ricettivo"}])
    c, _ = esegui_tutto(d, [], tmp_path, pdf=False)
    assert c.esito["rischio"]["rvita"].startswith("Ciii")
    assert {x["nome"] for x in c.esito["rischio"]["rvita_locali"]} == {"Uffici", "Camere"}


def test_minicodice_usa_valori_fissi(tmp_path):
    from antincendio_app.coordinatore import esegui_tutto
    d = dict(ragione_sociale="Z", indirizzo="V", datore_lavoro="A", rspp="B", superficie_mq=180, quota_min=0, quota_max=3, occupanti_lavoratori=8,
             tipologia_attivita=["ufficio"], sostanze_significative=False, lavorazioni_pericolose=False, qf_mj_m2=300)
    c, _ = esegui_tutto(d, [], tmp_path, pdf=False)
    dim = c.esito["dimensionamento"]
    assert dim["modo"] == "minicodice" and dim["les_max"] == 60 and dim["lcc_max"] == 15 and dim["uscite_min"] == 2   # sup > 150 m²


def test_percorso_scelto_dai_dati_non_assunto(tmp_path):
    from antincendio_app.coordinatore import esegui_tutto
    base = dict(ragione_sociale="Q", indirizzo="V", datore_lavoro="A", rspp="B", superficie_mq=300, occupanti_lavoratori=5, tipologia_attivita=["ufficio"], qf_mj_m2=300)
    # dati mancanti (quote, sostanze, lavorazioni): nessun requisito contraddetto -> Minicodice PROVVISORIO, non «non basso» di default
    c, _ = esegui_tutto(base, [], tmp_path / "a", pdf=False)
    n = c.esito["normativo"]
    assert n["template"] == "minicodice" and n["provvisoria"] is True and set(n["requisiti_ignoti"]) == {"C.c", "C.e", "C.f"}
    # con i dati completi la scelta diventa definitiva
    c, _ = esegui_tutto(dict(base, quota_min=0, quota_max=3, sostanze_significative=False, lavorazioni_pericolose=False), [], tmp_path / "b", pdf=False)
    assert c.esito["normativo"]["template"] == "minicodice" and c.esito["normativo"]["provvisoria"] is False
    # un requisito contraddetto -> Codice integrale (definitivo)
    c, _ = esegui_tutto(dict(base, sostanze_significative=True), [], tmp_path / "c", pdf=False)
    assert c.esito["normativo"]["template"] == "codice" and c.esito["normativo"]["requisiti_falliti"] == ["C.e"]
    # attivita soggetta -> raccordo
    c, _ = esegui_tutto(dict(base, tipologia_attivita=["commercio"], superficie_mq=600, occupanti_terzi=30), [], tmp_path / "d", pdf=False)
    assert c.esito["normativo"]["template"] == "raccordo"
    # il documento provvisorio non afferma «tutti i requisiti soddisfatti»
    import docx
    c, res = esegui_tutto(base, [], tmp_path / "e", pdf=False)
    d = docx.Document(str(tmp_path / "e" / res["documenti"][0]["file"]))
    t = "\n".join(p.text for tb in d.tables for r in tb.rows for cl in r.cells for p in cl.paragraphs)
    assert "CONCLUSIONE PROVVISORIA" in t and "CONCLUSIONE: tutti i requisiti" not in t


def test_confronto_versioni_e_config(tmp_path, monkeypatch):
    from antincendio_app import manutenzione as M
    assert M._vtuple("0.10.0") > M._vtuple("0.9.5") and M._vtuple("1.0") > M._vtuple("0.99.9")
    monkeypatch.setattr(M, "versione_remota", lambda: "99.0.0")
    assert M.controlla_aggiornamenti()["disponibile"] is True
    monkeypatch.setattr(M, "versione_remota", lambda: "0.0.1")
    assert M.controlla_aggiornamenti()["disponibile"] is False
    monkeypatch.setattr(M, "CONFIG", tmp_path / "config.json")
    M.salva_impostazioni({"auto_aggiorna": True, "altro": 1})
    assert M.leggi_impostazioni() == {"auto_aggiorna": True}


def test_aggiornamento_zip_preserva_dati_e_template(tmp_path, monkeypatch):
    import zipfile
    from antincendio_app import manutenzione as M
    monkeypatch.setattr(M, "ROOT", tmp_path / "app")
    (tmp_path / "app" / "templates").mkdir(parents=True); (tmp_path / "app" / "templates" / "t.docx").write_bytes(b"MIO")
    (tmp_path / "app" / "data").mkdir(); (tmp_path / "app" / "data" / "x.json").write_text("dati")
    z = tmp_path / "n.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("App-anti/antincendio_app/__init__.py", '__version__ = "9.9.9"\n')
        zf.writestr("App-anti/templates/t.docx", b"NUOVO")
        zf.writestr("App-anti/data/x.json", "sovrascritto?")
    r = M.applica_zip(z)
    assert r["versione_nuova"] == "9.9.9" and r["template_nuovi_da_confrontare"] == ["t.docx"]
    assert (tmp_path / "app" / "templates" / "t.docx").read_bytes() == b"MIO"
    assert (tmp_path / "app" / "templates" / "_nuovi" / "t.docx").read_bytes() == b"NUOVO"
    assert (tmp_path / "app" / "data" / "x.json").read_text() == "dati"
    assert "9.9.9" in (tmp_path / "app" / "antincendio_app" / "__init__.py").read_text()
