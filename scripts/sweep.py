"""Giro di prove automatico: genera molti scenari sintetici, esegue l'intera pipeline e controlla invarianti.
Uso: python scripts/sweep.py [N] [seed]   → stampa le anomalie trovate (da correggere)."""
import random, sys, tempfile, traceback, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import docx
from antincendio_app.coordinatore import esegui_tutto
from antincendio_app import checklist as CL, regole as R
from antincendio_app.agents import controllo

TIPI = ["ufficio", "commercio", "ristorazione", "officina/produzione", "deposito/magazzino", "ricettivo", "sanitario", "scuola", "spettacolo/intrattenimento", "autorimessa"]


def scenario(rng):
    tip = rng.sample(TIPI, rng.choice([1, 1, 1, 2]))
    d = dict(ragione_sociale=f"Azienda {rng.randint(1, 999)} S.r.l.", indirizzo="Via Prova 1 — 00100 Roma (RM)", datore_lavoro="Mario Rossi",
             rspp="Sara Santarelli", rls="Anna Bianchi", tipologia_attivita=tip, attivita_descrizione="Attività di prova")
    d["superficie_mq"] = rng.choice([60, 180, 350, 650, 990, 1010, 1800, 5200])
    d["occupanti_lavoratori"] = rng.choice([1, 4, 12, 40, 95, 130, 420])
    d["occupanti_terzi"] = rng.choice([0, 0, 8, 40, 120])
    d["aperta_pubblico"] = d["occupanti_terzi"] > 0
    d["quota_min"] = rng.choice([0, 0, -3, -8]); d["quota_max"] = rng.choice([3, 8, 14, 26, 40])
    d["piani_n"] = rng.choice([1, 1, 2, 4])
    if rng.random() < .7: d["sostanze_significative"] = rng.random() < .35
    if rng.random() < .7: d["lavorazioni_pericolose"] = rng.random() < .3
    if rng.random() < .5: d["qf_mj_m2"] = rng.choice([80, 250, 450, 700, 950, 1300, 2600])
    if "ricettivo" in tip: d["posti_letto"] = rng.choice([8, 24, 60, 120])
    if "sanitario" in tip and rng.random() < .5: d["posti_letto"] = rng.choice([10, 40])
    if "ristorazione" in tip and rng.random() < .5:
        d.update(eventi_intrattenimento=True, pista_ballo=rng.random() < .4, ingresso_pagamento=rng.random() < .3, eventi_solo_aperto=rng.random() < .15)
    if rng.random() < .4:
        d["locali"] = [{"nome": "Locale A", "mq": d["superficie_mq"] * .6, "tipo": tip[0]}, {"nome": "Locale B", "mq": d["superficie_mq"] * .4, "tipo": rng.choice(TIPI)}]
    if rng.random() < .3: d["n_uscite"] = rng.choice([1, 2, 3]); d["les_misurata"] = rng.choice([15, 45, 90]); d["lcc_misurata"] = rng.choice([5, 18, 40]); d["largh_installata_mm"] = rng.choice([800, 1200, 2400]); d["n_estintori"] = rng.choice([1, 3, 9])
    if rng.random() < .3: d.update(litri_infiammabili=rng.choice([20, 80, 150, 1500]), kg_carta=rng.choice([100, 6000]), kw_termico=rng.choice([60, 200]))
    if rng.random() < .3: d.update(soggetta_dpr151=rng.random() < .5, categoria_dpr151="B", scadenza_rinnovo=rng.choice(["10/03/2020", "01/01/2031"]), tipo_procedimento="SCIA art. 4")
    if rng.random() < .2: d.update(vincolata=rng.random() < .5, strategica=rng.random() < .5)
    return d


def controlla(d, c, res, out):
    p = []
    n, r = c.esito["normativo"], c.esito["rischio"]
    req = n["requisiti_allegato_I"]
    if (n["ramo"] == "MINICODICE") != (all(x["ok"] is True for x in req) and not d.get("rt_settore") and not d.get("rtv_applicabile")):
        p.append("ramo/requisiti incoerenti")
    if not re.fullmatch(r"(A|B|Ci{1,3}|D|E)[1-4]", r["rvita"]): p.append(f"Rvita malformato {r['rvita']}")
    m = re.fullmatch(r"(A|B|Ci{1,3}|D|E)([1-4])", r["rvita"])
    if m and not R.rvita_ammesso(m.group(1), int(m.group(2))): p.append(f"Rvita non ammesso {r['rvita']}")
    if r.get("qfd") and not (r["qfd"]["qfd"] > 0): p.append("qfd non positivo")
    if n["soggetta_dpr151"] and n["template"] != "raccordo": p.append("soggetta ma template non raccordo")
    if (not n["soggetta_dpr151"]) and n["ramo"] in ("A", "B") : p.append("ramo A/B non soggetta")
    for doc in res["documenti"]:
        if doc["verifica"]: p.append(f"verifica doc: {doc['verifica']}")
        t = docx.Document(str(out / doc["file"]))
        txt = "\n".join(x.text for x in t.paragraphs)
        if "GUIDA ALLA COMPILAZIONE" in txt: p.append("guide residue")
        if d["ragione_sociale"].upper() not in txt.upper() and doc["tipo"] == "VRI": p.append("ragione sociale assente")
    if res["stato"] != "BOZZA": p.append("stato non BOZZA senza sopralluogo")
    voci = CL.estrai(docx.Document(str(Path(__file__).resolve().parents[1] / "templates" / {"minicodice": "VRI_Minicodice_template_2.docx", "codice": "VRI_Codice_Integrale_RTO_completa_2.docx", "raccordo": "VRI_Raccordo_CPI_template.docx"}[n["template"]])), n["template"])
    if len(voci) < 20: p.append("check-list troppo corta")
    return p


def main(N=40, seed=1):
    rng = random.Random(seed); bad = 0
    for i in range(N):
        d = scenario(rng)
        out = Path(tempfile.mkdtemp())
        try:
            c, res = esegui_tutto(d, [], out, pdf=False)
            pr = controlla(d, c, res, out)
        except Exception:
            pr = ["ECCEZIONE: " + traceback.format_exc().splitlines()[-1]]
            c = None
        if pr:
            bad += 1
            print(f"[{i}] {d.get('tipologia_attivita')} sup={d.get('superficie_mq')} occ={d.get('occupanti_lavoratori')}+{d.get('occupanti_terzi')}: {pr}")
            if pr[0].startswith("ECCEZIONE"):
                print("    ", {k: v for k, v in d.items() if k not in ('ragione_sociale', 'indirizzo', 'datore_lavoro', 'rspp', 'rls')})
    print(f"\n{N} scenari, {bad} con anomalie")
    return bad


if __name__ == "__main__":
    sys.exit(1 if main(int(sys.argv[1]) if len(sys.argv) > 1 else 40, int(sys.argv[2]) if len(sys.argv) > 2 else 1) else 0)
